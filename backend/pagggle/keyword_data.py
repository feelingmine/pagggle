"""Offline keyword tables and explicit SERP snapshots; never queries a search API."""

import base64
import csv
import io
import json
import re
from datetime import date, datetime
from urllib.parse import urlsplit, urlunsplit
from zipfile import BadZipFile, ZipFile


COLUMNS = {
    "keyword": "keyword", "search volume": "volume", "volume": "volume",
    "keyword difficulty": "kd", "keyword difficulty %": "kd", "kd": "kd",
    "serp results": "serp", "serp": "serp", "market": "market", "language": "language",
    "source": "source", "data_date": "data_date", "serp_date": "serp_date",
    "serp_source": "serp_source", "device": "device",
}


def read_table(text, format):
    if format == "csv":
        try:
            rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff")), strict=True))
        except csv.Error:
            raise ValueError("CSV 格式无效，请检查引号和分隔符") from None
    else:
        from openpyxl import load_workbook

        try:
            raw = base64.b64decode(text, validate=True)
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 20_000_000:
                    raise ValueError("XLSX 解压后超过 20 MB，请拆分文件")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.sheetnames) != 1:
                    raise ValueError("请提供仅含一个工作表的 XLSX，避免遗漏其他工作表")
                rows = []
                for row in workbook.active.iter_rows():
                    if len(rows) > 1000 or len(row) > 50:
                        raise ValueError("每次支持最多 1000 条、50 列，请拆分文件")
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("XLSX 含公式，请先复制为值后导入")
                    values = [cell.value for cell in row]
                    rows.append([v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else "" if v is None else str(v) for v in values])
            finally:
                workbook.close()
        except (BadZipFile, KeyError, OSError, TypeError):
            raise ValueError("XLSX 文件无效，请检查文件格式") from None
    if not rows:
        raise ValueError("表格为空")
    headers = [COLUMNS.get(str(h).strip().casefold()) for h in rows[0]]
    if None in headers or "keyword" not in headers or len(set(headers)) != len(headers):
        raise ValueError("需要 keyword / Keyword 列；支持搜索量、KD、SERP Results、market、language、source、data_date、serp_date、serp_source、device；请移除其他列或重复列")
    result = []
    for row in rows[1:]:
        if not any(str(v).strip() for v in row):
            continue
        record = {h: row[i] if i < len(row) else None for i, h in enumerate(headers)}
        if len(row) > len(headers):
            record[None] = row[len(headers):]
        result.append(record)
    return result


def parse_serp(value):
    """Preserve rank order, up to the top 10; compare full URLs, not domains."""
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        if value.startswith("["):
            try:
                value = json.loads(value)
            except ValueError:
                raise ValueError("SERP JSON 列表无效") from None
        else:
            value = re.split(r"[,\r\n]+", value)
    if not isinstance(value, list) or not all(isinstance(url, str) for url in value):
        raise ValueError("SERP 需要网址列表，不能填排名、域名指标或 SERP Features")
    urls = []
    for url in value:
        url = url.strip()
        if not url:
            continue
        if "://" not in url:
            url = "https://" + url
        try:
            parts = urlsplit(url)
            if parts.scheme not in {"http", "https"} or not parts.hostname or "." not in parts.hostname or parts.username or parts.password or any(c.isspace() for c in url):
                raise ValueError
            port = parts.port
            host = parts.hostname.lower().encode("idna").decode()
            if port and port != (443 if parts.scheme == "https" else 80):
                host += f":{port}"
            normalized = urlunsplit((parts.scheme.lower(), host, parts.path or "/", parts.query, ""))
        except (ValueError, UnicodeError):
            raise ValueError("SERP 中包含无效网址") from None
        if normalized not in urls:
            urls.append(normalized)
    if len(urls) > 10:
        raise ValueError("每个关键词最多导入前 10 个自然结果 URL，请按排名整理；不会自动截断")
    return urls or None
