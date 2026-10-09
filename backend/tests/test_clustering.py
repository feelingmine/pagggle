import base64
import csv
import io

import pytest
from openpyxl import Workbook

from pagggle.clustering import cluster_keywords
from pagggle.config import Settings
from pagggle.intake import parse_records, fingerprint
from pagggle.keyword_data import parse_serp, read_table


def settings(**values):
    return Settings(BASE_URL='https://provider.example/v1', model='test', api_key='test', clustering_serp_threshold=2, **values)


def record(id, serp, volume=100, kd=20, **values):
    return dict(id=id, kind='keyword', status='pending', original=id, normalized=id, serp=serp, volume=volume, kd=kd,
                market='US', language='en', serp_source='fixture', serp_date='2026-10-08', device='desktop', **values)


def run(rows, **kwargs):
    return cluster_keywords(settings(), rows, encoder=lambda texts: [[1, 0]] * len(texts), **kwargs)


def test_semantic_silos_precede_serp_split_and_exactly_one_primary():
    urls = ['https://example.com/a', 'https://example.com/b']
    rows = [record('gasket', urls, 300, 40), record('seal', urls, 300, 20), record('seo', urls, 200)]
    vectors = {'gasket': [1, 0], 'seal': [1, 0.05], 'seo': [-1, 0]}
    result = cluster_keywords(settings(), rows, encoder=lambda texts: [vectors[t] for t in texts])
    assert result['summary']['silos'] == 2
    assert result['summary']['groups'] == 2
    for group in result['groups']:
        assert sum(m['keyword_role'] == 'Primary' for m in group['members']) == 1
    pair = next(g for g in result['groups'] if len(g['members']) == 2)
    assert pair['members'][0]['keyword'] == 'seal'
    assert pair['pair_evidence'][0]['overlap'] == 2
    assert pair['status'] == 'needs_review'


def test_pooled_serp_bridge_is_not_approved_and_input_order_is_stable():
    rows = [record('a', ['a', 'b']), record('b', ['a', 'b', 'c', 'd']), record('c', ['c', 'd'])]
    result = run(rows)
    group = result['groups'][0]
    assert len(group['members']) == 3 and group['chain_overlap']
    assert group['status'] == 'needs_review'
    assert min(p['overlap'] for p in group['pair_evidence']) == 0
    assert run(rows[::-1])['groups'] == result['groups']


def test_missing_serp_unknown_metrics_and_distinct_snapshot_contexts():
    rows = [record('a', None, None, None), record('b', None, 0, 0)]
    result = run(rows)
    assert len(result['groups']) == 2
    assert all(g['status'] == 'needs_evidence' for g in result['groups'])
    assert next(g for g in result['groups'] if g['members'][0]['keyword'] == 'a')['members'][0]['volume'] is None
    a, b = record('a', ['x', 'y']), record('b', ['x', 'y'])
    for key, value in [('market', 'DE'), ('language', 'de'), ('serp_date', '2026-10-09'), ('device', 'mobile'), ('serp_source', 'different')]:
        b = dict(a, id='b', original='b', normalized='b', **{key: value})
        assert run([a, b])['summary']['groups'] == 2
    b = dict(a, id='b', original='b', normalized='b', volume=None)
    group = run([a, b])['groups'][0]
    assert group['members'][0]['keyword'] == 'a' and group['primary_provisional']


def test_singleton_cancellation_and_invalid_vectors():
    rows = [record('a', None)]
    assert run(rows)['summary']['keywords'] == 1
    with pytest.raises(ValueError, match='取消'):
        run(rows, cancelled=lambda: True)
    for vectors in ([[0, 0]], [[float('nan'), 1]], [[1, 0], [1, 0]]):
        with pytest.raises(ValueError, match='向量'):
            cluster_keywords(settings(), rows, encoder=lambda _: vectors)
    with pytest.raises(ValueError):
        run([])


def test_serp_normalization_full_urls_unknown_invalid_and_top_ten():
    assert parse_serp(' HTTPS://EXAMPLE.COM:443/a#one,https://example.com/a\nexample.com/b?q=1 ') == ['https://example.com/a', 'https://example.com/b?q=1']
    assert parse_serp('') is None and parse_serp('[]') is None
    assert parse_serp('["https://example.com/a","https://example.com/b"]') == ['https://example.com/a', 'https://example.com/b']
    for value in ('javascript:alert(1)', '[1]', 'https://user:password@example.com', ','.join(f'https://example.com/{i}' for i in range(11))):
        with pytest.raises(ValueError):
            parse_serp(value)


def test_semrush_headers_csv_xlsx_equivalence_and_serp_fingerprint():
    headers = ['Keyword', 'Search Volume', 'Keyword Difficulty', 'SERP Results', 'market', 'language', 'source', 'data_date', 'device']
    rows = [['silicone mold', '', 12.5, 'https://a.example/\nhttps://b.example/', 'US', 'en', 'Semrush export', '2026-10-08', 'desktop'], ['bad keyword', -1, 101, '', '', '', '', '', '']]
    out = io.StringIO(); writer = csv.writer(out); writer.writerow(headers); writer.writerows(rows)
    data = dict(kind='keyword', text=out.getvalue(), format='csv')
    csv_records = parse_records(data)
    assert csv_records[0]['volume'] is None and csv_records[0]['kd'] == 12.5
    assert csv_records[0]['serp_date'] == '2026-10-08'
    assert csv_records[1]['status'] == 'failed'
    wb = Workbook(); sheet = wb.active; sheet.append(headers)
    for row in rows: sheet.append(row)
    buf = io.BytesIO(); wb.save(buf)
    xlsx_records = parse_records(dict(data, text=base64.b64encode(buf.getvalue()).decode(), format='xlsx'))
    assert xlsx_records == csv_records
    assert fingerprint(dict(csv_records[0], serp_date='2026-10-09')) != csv_records[0]['fingerprint']
    wb.create_sheet('another'); buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(ValueError, match='一个工作表'):
        read_table(base64.b64encode(buf.getvalue()).decode(), 'xlsx')


def test_extra_headers_preserved_without_treating_features_as_serp_urls():
    data = parse_records(dict(kind='keyword', format='csv', text='Keyword,SERP Features\na,video'))
    assert data[0]['raw']['SERP Features'] == 'video'
    assert data[0]['serp'] is None
    for text in ('Keyword,keyword\na,a', 'Keyword,Note,note\na,b,c', 'Keyword,\na,b'):
        with pytest.raises(ValueError): read_table(text, 'csv')
    wb = Workbook(); wb.active.append(['Keyword']); wb.active.append(['=1+1'])
    buf = io.BytesIO(); wb.save(buf)
    with pytest.raises(ValueError, match='公式'):
        read_table(base64.b64encode(buf.getvalue()).decode(), 'xlsx')
